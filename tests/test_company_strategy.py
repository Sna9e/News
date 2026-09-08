import copy
import datetime as dt
import json
import sys
import unittest
from pathlib import Path
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from tools import company_strategy as cs
from tools import finance_engine as fe
from tools.company_query_packs import DEFAULT_COMPANY_TOPICS, get_company_query_pack
from tools.export_ppt import generate_ppt
from pptx import Presentation

NOW = dt.datetime(2026, 9, 7, tzinfo=dt.timezone.utc)
URL = "https://www.microsoft.com/DEMO-only"
QUOTE = "DEMO: Fourth quarter ended June 30, 2026 capital expenditures including finance leases were $30 billion."


def fixture():
    source = {"url": URL, "content": QUOTE, "company": "Microsoft", "kind": "capex", "method": "jina", "published_date": "2026-08-01"}
    claim = cs.EvidenceClaim(text="DEMO资本支出$30 billion", source_url=URL, evidence_quote=QUOTE)
    draft = cs.StrategyOverview(capex=[cs.CapexRow(company="Microsoft", period="Fourth quarter", period_end="2026-06-30",
                                                basis="including finance leases", quarter=claim)])
    return draft, [source]


class StrategyTests(unittest.TestCase):
    def test_terms_are_annotated_once_without_mutating_source_text(self):
        from tools.export_ppt import _annotate_first_terms
        from pptx.util import Inches
        prs = Presentation()
        page = prs.slides.add_slide(prs.slide_layouts[6])
        frame = page.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(4)).text_frame
        first = frame.paragraphs[0]
        first.text = "CPO与HBM（高带宽存储器），采用Optical I/O。"
        second = frame.add_paragraph()
        second.text = "CPO、HBM、Optical I/O再次出现。"
        _annotate_first_terms(prs)
        self.assertIn("CPO（共封装光学）", first.text)
        self.assertIn("Optical I/O（光学输入/输出）", first.text)
        self.assertEqual(first.text.count("高带宽存储器"), 1)
        self.assertNotIn("（", second.text)

    def test_fake_ai_uses_existing_search_and_crawler_and_passes_fields(self):
        draft, sources = fixture()
        class FakeAI:
            valid = True
            calls = 0
            def analyze_structural(self, prompt, schema):
                self.calls += 1
                assert "including finance leases" in prompt
                return draft
        ai = FakeAI()
        queries = []
        def search(query, sites, time_flag, **kwargs):
            queries.append(query)
            return [{**sources[0], "title": "DEMO Fourth quarter results"}] if "Microsoft" in query else []
        def fetch(url, **kwargs):
            return {"text": QUOTE, "method": "jina"}
        output = cs.build_strategy_overview(ai, {}, now=NOW, search_fn=search, fetch_fn=fetch)
        self.assertEqual(ai.calls, 1)
        self.assertEqual(len(queries), 4)
        self.assertEqual(output["capex"][0]["quarter"]["text"], "DEMO资本支出$30 billion")
        sections = [{"topic": "微软", "report_style": "company_tracking", "finance": {}},
                    {"topic": "SpaceX", "report_style": "company_tracking", "finance": {"catalysts": {"style": "原内容"}}},
                    {"topic": "其他频道", "report_style": "consumer_daily", "finance": {}}]
        cs.apply_strategy_to_sections(sections, output)
        self.assertIn("$30 billion", sections[0]["finance"]["catalysts"]["earnings"])
        self.assertEqual(sections[1]["finance"]["catalysts"]["style"], "原内容")
        self.assertNotIn("strategy_overview", sections[2])

    def test_company_pack_keeps_old_topics_and_adds_microsoft(self):
        self.assertTrue({"Microsoft", "OpenAI", "Anthropic", "SpaceX", "特朗普"}.issubset(DEFAULT_COMPANY_TOPICS))
        self.assertIn("Maia", get_company_query_pack("微软")["keywords"])
        self.assertIn("TPU", get_company_query_pack("Google")["priority_terms"])

    def test_official_capex_requires_exact_evidence_and_period(self):
        draft, sources = fixture()
        self.assertTrue(cs.validate_overview(draft, sources, NOW)["capex"][0]["quarter"])
        for mutation in ("amount", "period", "company", "basis", "snippet", "future"):
            d, s = copy.deepcopy(draft), copy.deepcopy(sources)
            if mutation == "amount": d.capex[0].quarter.text = "$999 billion"
            if mutation == "period": d.capex[0].period = "Full year"
            if mutation == "company": s[0]["company"] = "Meta"
            if mutation == "basis": d.capex[0].basis = "cash only"
            if mutation == "snippet": s[0]["method"] = "search_snippet"
            if mutation == "future": d.capex[0].period_end = "2027-01-01"
            self.assertFalse(cs.validate_overview(d, s, NOW)["capex"][0]["quarter"], mutation)

    def test_source_boundary_and_unverified_numbers(self):
        self.assertEqual(cs.source_level("https://microsoft.com.evil.test/"), "D")
        self.assertEqual(cs.source_level("https://news.microsoft.com/article"), "A")
        self.assertEqual(cs.source_level("https://podcasts.apple.com/user-podcast"), "D")
        self.assertEqual(cs.source_level("https://learn.microsoft.com/answers/questions/123"), "D")
        d, s = fixture()
        d.capex[0].quarter.text = "DEMO $30 million"
        self.assertFalse(cs.validate_overview(d, s, NOW)["capex"][0]["quarter"])
        d, s = fixture()
        d.capex[0].quarter.evidence_quote = "Capital expenditures were $30 billion but this is not in source"
        self.assertFalse(cs.validate_overview(d, s, NOW)["capex"][0]["quarter"])

    def test_no_direction_from_missing_or_weak_evidence(self):
        excerpt = "DEMO Microsoft announced deployment of a new GPU AI compute cluster for customers."
        claim = cs.EvidenceClaim(text="DEMO新GPU集群开始部署", source_url=URL, evidence_quote=excerpt)
        source = {"url": URL, "content": excerpt, "company": "Microsoft", "kind": "strategy"}
        draft = cs.StrategyOverview(trends=[cs.TrendRow(name="AI算力", direction="增强", evidence=[claim])])
        self.assertEqual(cs.validate_overview(draft, [source], NOW)["trends"][0]["arrow"], "↑")
        draft.trends[0].direction = "持平"
        self.assertFalse(cs.validate_overview(draft, [source], NOW)["trends"])
        draft.trends[0].direction = "增强"
        source["url"] = claim.source_url = "https://ordinary-blog.test/article"
        self.assertFalse(cs.validate_overview(draft, [source], NOW)["trends"])

    def test_secondary_information_and_rumors_guard(self):
        from agents.deep_analyst import NewsItem, enforce_company_evidence
        news = NewsItem(title="DEMO技术更新", source="媒体", date_check="2026-09-07", url="https://example.com/news",
                        summary="【事件核心】技术更新。\n【深度细节/数据支撑】报道客户部署。\n【行业深远影响】重塑行业。", importance=5)
        raw = [{"url": news.url, "title": "DEMO技术更新", "content": "客户部署了新服务器"}]
        kept, _ = enforce_company_evidence([news], raw)
        self.assertIn("媒体报道，尚未获得官方确认", kept[0].summary)
        self.assertNotIn("重塑行业", kept[0].summary)
        self.assertEqual(kept[0].importance, 3)
        raw[0]["content"] = "unconfirmed rumor"
        self.assertFalse(enforce_company_evidence([news], raw)[0])

    def test_tencent_us_fields_do_not_become_pb_or_year_range(self):
        parts = [""] * 52
        for index, value in {2: "AAPL.OQ", 3: "200", 4: "199", 5: "198", 32: "0.5", 33: "201", 34: "197", 39: "30", 45: "30000", 46: "Apple Inc."}.items():
            parts[index] = value
        class Response:
            content = "~".join(parts).encode("gbk")
        with patch.object(fe, "_request_response", return_value=Response()):
            result = fe.fetch_tencent_quote_metrics("AAPL")
            self.assertEqual(result["pe"], "30")
            self.assertNotIn("pb", result)
            self.assertNotIn("low_52w", result)
        self.assertIsNone(fe._safe_float(float("inf")))

    def test_secondary_quote_never_replaces_primary_price(self):
        from test_finance_engine import _yahoo_payload, FakeResponse
        payload = _yahoo_payload("AAPL")
        meta = payload["chart"]["result"][0]["meta"]
        for key in ("trailingPE", "priceToBook"): meta.pop(key)
        with patch.object(fe, "_request_response", return_value=FakeResponse(payload)), \
             patch.object(fe, "fetch_yahoo_quote_metrics", return_value={}), \
             patch.object(fe, "fetch_tencent_quote_metrics", return_value={"pe": 30, "current_price": 1, "currency": "CNY", "high_52w": 2}), \
             patch.object(fe, "generate_pro_kline_chart", return_value=None):
            result = fe.fetch_from_yahoo_chart("AAPL")
        self.assertEqual(result["current_price"], 136.97)
        self.assertEqual(result["currency"], "USD")
        self.assertEqual(result["range_52w"], "128.00 - 151.00")

    def test_stub_overview_adds_exactly_two_pages_and_preserves_existing_slides(self):
        draft, sources = fixture()
        overview = cs.validate_overview(draft, sources, NOW)
        directory = ROOT / "validation_outputs/strategy_2026_09_07"
        directory.mkdir(parents=True, exist_ok=True)
        news = {"title": "DEMO：技术战略排版测试", "date_check": "2026-09-07", "source": "DEMO", "url": URL,
                "summary": "【事件核心】\n这是一条仅用于本地排版验证的演示新闻。\n【深度细节/数据支撑】\n演示数据不代表公司实际公告。\n【行业深远影响】\n不得将演示数据用于投资判断。"}
        sections = [{"topic": "Microsoft", "data": [news], "report_style": "company_tracking", "finance": {}}]
        baseline = Presentation(generate_ppt(sections, [], str(directory / "baseline"), "DEMO"))
        sections[0]["strategy_overview"] = overview
        generated = Presentation(generate_ppt(sections, [], str(directory / "strategy_stub"), "DEMO"))
        self.assertEqual(len(generated.slides), len(baseline.slides) + 2)
        self.assertEqual(generated.slides[-1]._element.xml, baseline.slides[-1]._element.xml)
        texts = " ".join(shape.text for page in generated.slides for shape in page.shapes if shape.has_text_frame)
        self.assertEqual(texts.count("全球科技趋势与头部企业战略动态"), 2)
        links = [run.hyperlink.address for page in generated.slides for shape in page.shapes if shape.has_table
                 for row in shape.table.rows for cell in row.cells for p in cell.text_frame.paragraphs for run in p.runs]
        self.assertIn(URL, links)
        for page in generated.slides:
            for shape in page.shapes:
                self.assertLessEqual(shape.left + shape.width, generated.slide_width + 100)
                self.assertLessEqual(shape.top + shape.height, generated.slide_height + 100)
        summary = {"added_pages": 2, "original_detail_slide_xml_unchanged": True, "official_link_preserved": True,
                   "note": "DEMO fixture, not live model output"}
        (directory / "stub_checks.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")

    def test_filled_layout_and_unchanged_market_chart(self):
        from test_finance_ppt_output import _finance_payload
        finance = _finance_payload()
        directory = ROOT / "validation_outputs/strategy_2026_09_07"
        directory.mkdir(parents=True, exist_ok=True)
        claim = lambda text: {"text": text, "source_url": URL, "source_level": "A"}
        overview = {"capex": [], "companies": [], "trends": []}
        for company, config in cs.CONFIG["companies"].items():
            if config.get("ir_domains"):
                overview["capex"].append({"company": company, "period": "DEMO Q2 2026", "basis": "含融资租赁",
                    "quarter": claim("DEMO $30 billion"), "yoy": claim("DEMO 20%"), "guidance": claim("DEMO $120 billion"),
                    "revision": claim("DEMO 上调"), "destination": claim("DEMO：AI服务器、数据中心网络与电力基础设施")})
            overview["companies"].append({"company": company, "action": claim("DEMO：公司新增AI服务器部署，网络平台同步升级，客户正在开展适配验证"),
                "direction": "AI基础设施", "key_data": claim("DEMO：客户验证阶段"),
                "meaning": claim("若客户验证持续推进，可能增加服务器与高速互连的适配需求"),
                "earnings": claim("DEMO：公司调整资本投入方向，重点覆盖服务器、网络和数据中心供电"),
                "transmission": claim("若部署规模扩大，可能传导至芯片、内存、封装、高速互连和电源"),
                "policy": claim("DEMO：监管变化可能影响芯片供应范围和数据中心项目审批节奏")})
        overview["trends"] = [{"name": name, "arrow": "↑↑", "evidence": [claim("DEMO：新增客户验证与产品部署，相关接口和供应链适配需求增加")]} for name in cs.CONFIG["trends"]]
        sections = [{"topic": "Microsoft", "report_style": "company_tracking", "finance": finance,
                     "data": [{"title": "DEMO：满载排版测试，不代表真实公告", "summary": "仅用于测试。", "source": "DEMO"}]}]
        cs.apply_strategy_to_sections(sections, overview)
        result = Presentation(generate_ppt(sections, [], str(directory / "strategy_filled_stub"), "DEMO"))
        financial = next(s for s in result.slides if any("量化面与事件催化" in sh.text for sh in s.shapes if sh.has_text_frame))
        picture = next(sh for sh in financial.shapes if sh.shape_type == 13)
        self.assertEqual(picture.left.inches, 4.5)
        self.assertEqual(picture.top.inches, 1.2)
        self.assertEqual(picture.width.inches, 5)
        self.assertEqual(picture.image.blob, Path(finance["chart_path"]).read_bytes())
        self.assertNotIn("市场风格轮动", " ".join(sh.text for sh in financial.shapes if sh.has_text_frame))


if __name__ == "__main__":
    unittest.main()
